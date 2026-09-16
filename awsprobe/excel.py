"""inventory.json から Excel 棚卸し表（.xlsx）を作る。

`docs/INVENTORY_SCHEMA.md` の各セクションと、コレクタの実装
（`awsprobe/collectors/*.py`）を突き合わせて読み、実際に収集されている
キーだけを見て埋める。契約（スキーマ文書）とコレクタ実装がずれている
箇所は **コレクタ実装を正** として扱う。

体裁は次のとおり
（濃紺ヘッダ＋白文字太字・縞模様・薄灰枠線・フリーズペイン・
オートフィルタ・折り返し・Yu Gothic フォント）。

厳守事項:
- **AWS API を呼ばない。boto3 を import しない。** 入力は inventory dict のみ。
- inventory のキーが欠けていても KeyError で落ちない（`.get()` を徹底する）。
- アクセスキー ID の全体・シークレット値・SSM パラメータ値は出力しない
  （inventory 側で既にマスク・除外済みのものをそのまま転記するだけであり、
  ここで追加のマスキングは行わないが、生成物にそれらのキーを新たに
  持ち込むことも絶対にしない）。
- 依存は openpyxl と標準ライブラリのみ。
"""
from __future__ import annotations

import datetime as _dt
import json
import re
import unicodedata
from typing import Any, Callable, Sequence

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

# ---------------------------------------------------------------------------
# 体裁定数（既存 xlsx 実物から採取した値に合わせる）
# ---------------------------------------------------------------------------

_FONT_NAME = "Yu Gothic"
_FONT_SIZE = 10

_HEADER_FILL_RGB = "FF1F3864"
_HEADER_FONT_RGB = "FFFFFFFF"
_STRIPE_FILL_RGB = "FFF2F5FA"
_BORDER_RGB = "FFBFBFBF"
_DATA_FONT_RGB = "FF000000"

_SUBTITLE_FILL_RGB = "FFD9E2F3"
_SUBTITLE_FONT_RGB = "FF1F3864"

#: Excel の「不正」「注意」相当の配色（未実施＝赤系／一部実施＝黄系／0.0.0.0/0＝赤系）
_RED_FILL_RGB = "FFFFC7CE"
_RED_FONT_RGB = "FF9C0006"
_YELLOW_FILL_RGB = "FFFFEB9C"
_YELLOW_FONT_RGB = "FF9C6500"

_THIN_SIDE = Side(style="thin", color=_BORDER_RGB)
_BORDER = Border(left=_THIN_SIDE, right=_THIN_SIDE, top=_THIN_SIDE, bottom=_THIN_SIDE)

_MIN_COL_WIDTH = 9.0
_MAX_COL_WIDTH = 60.0
_MAX_CELL_LEN = 30000  # Excel のセル文字数上限（32767）に対する安全マージン

_ILLEGAL_XML_RE = re.compile(r"[\x00-\x08\x0B\x0C\x0E-\x1F]")


def _header_font() -> Font:
    return Font(name=_FONT_NAME, size=_FONT_SIZE, bold=True, color=_HEADER_FONT_RGB)


def _header_fill() -> PatternFill:
    return PatternFill(patternType="solid", fgColor=_HEADER_FILL_RGB)


def _header_align() -> Alignment:
    return Alignment(horizontal="center", vertical="center", wrap_text=True)


def _data_font(color: str | None = None) -> Font:
    return Font(name=_FONT_NAME, size=_FONT_SIZE, color=color or _DATA_FONT_RGB)


def _data_align() -> Alignment:
    return Alignment(horizontal="left", vertical="top", wrap_text=True)


def _subtitle_font() -> Font:
    return Font(name=_FONT_NAME, size=11, bold=True, color=_SUBTITLE_FONT_RGB)


def _subtitle_fill() -> PatternFill:
    return PatternFill(patternType="solid", fgColor=_SUBTITLE_FILL_RGB)


# ---------------------------------------------------------------------------
# セル値の整形
# ---------------------------------------------------------------------------


def _sanitize_text(text: str) -> str:
    """XML として不正な制御文字を落とし、長すぎる文字列を切り詰める。"""
    text = _ILLEGAL_XML_RE.sub("", text)
    if len(text) > _MAX_CELL_LEN:
        text = text[:_MAX_CELL_LEN] + "…(省略)"
    return text


def _cell_value(value: Any) -> Any:
    """任意の値を openpyxl セルに書ける形に変換する。

    - None / 空リスト / 空 dict → None（空セル）
    - bool → "はい" / "いいえ"
    - int / float → そのまま（数値として書く）
    - list / tuple / set → "; " 区切りの文字列
    - dict → JSON 文字列（機微情報は inventory 側で既に除外済みの前提）
    - それ以外 → str() し、不正文字を除去
    """
    if value is None:
        return None
    if isinstance(value, bool):
        return "はい" if value else "いいえ"
    if isinstance(value, (int, float)):
        return value
    if isinstance(value, (list, tuple, set)):
        parts = [str(_cell_value(v)) for v in value if v is not None and v != ""]
        return _sanitize_text("; ".join(parts)) if parts else None
    if isinstance(value, dict):
        if not value:
            return None
        try:
            return _sanitize_text(json.dumps(value, ensure_ascii=False, default=str))
        except (TypeError, ValueError):
            return _sanitize_text(str(value))
    return _sanitize_text(str(value))


def _display_width(value: Any) -> float:
    """日本語（全角）を考慮した表示幅。列幅の自動調整に使う。"""
    text = "" if value is None else str(value)
    width = 0.0
    for ch in text:
        east = unicodedata.east_asian_width(ch)
        width += 1.9 if east in ("W", "F", "A") else 1.0
    return width


# ---------------------------------------------------------------------------
# 低レベルの表描画
# ---------------------------------------------------------------------------


def _write_header_row(ws: Worksheet, row: int, headers: Sequence[str]) -> None:
    for col, text in enumerate(headers, start=1):
        cell = ws.cell(row=row, column=col, value=text)
        cell.font = _header_font()
        cell.fill = _header_fill()
        cell.alignment = _header_align()
        cell.border = _BORDER


def _write_data_row(
    ws: Worksheet,
    row: int,
    values: Sequence[Any],
    *,
    stripe: bool = False,
    fill_rgb: str | None = None,
    font_rgb: str | None = None,
) -> None:
    for col, value in enumerate(values, start=1):
        cell = ws.cell(row=row, column=col, value=_cell_value(value))
        cell.font = _data_font(font_rgb)
        cell.alignment = _data_align()
        cell.border = _BORDER
        bg = fill_rgb or (_STRIPE_FILL_RGB if stripe else None)
        if bg:
            cell.fill = PatternFill(patternType="solid", fgColor=bg)


def _write_title_row(ws: Worksheet, row: int, text: str, ncols: int) -> None:
    """複数表を1シートに積む際の見出し帯（表と表の間の区切り）。"""
    ncols = max(ncols, 1)
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=ncols)
    cell = ws.cell(row=row, column=1, value=text)
    cell.font = _subtitle_font()
    cell.fill = _subtitle_fill()
    cell.alignment = Alignment(horizontal="left", vertical="center")
    for col in range(1, ncols + 1):
        ws.cell(row=row, column=col).fill = _subtitle_fill()
        ws.cell(row=row, column=col).border = _BORDER


def _write_table(
    ws: Worksheet,
    start_row: int,
    headers: Sequence[str],
    rows: Sequence[Sequence[Any]],
    *,
    row_fill: Callable[[Sequence[Any]], tuple[str, str] | None] | None = None,
) -> int:
    """headers を start_row に、rows をその下に書く。次の空き行番号を返す。"""
    _write_header_row(ws, start_row, headers)
    for i, row_values in enumerate(rows):
        row_idx = start_row + 1 + i
        fill_rgb = font_rgb = None
        if row_fill is not None:
            picked = row_fill(row_values)
            if picked:
                fill_rgb, font_rgb = picked
        _write_data_row(
            ws, row_idx, row_values,
            stripe=(i % 2 == 1), fill_rgb=fill_rgb, font_rgb=font_rgb,
        )
    return start_row + 1 + len(rows)


def _autosize(ws: Worksheet, blocks: Sequence[tuple[Sequence[str], Sequence[Sequence[Any]]]]) -> None:
    """複数ブロック分の headers/rows から列幅を決める。"""
    widths: dict[int, float] = {}
    for headers, rows in blocks:
        for col, h in enumerate(headers, start=1):
            widths[col] = max(widths.get(col, _MIN_COL_WIDTH), _display_width(h) + 4)
        for row_values in rows:
            for col, v in enumerate(row_values, start=1):
                text = _cell_value(v)
                w = _display_width(text) + 4
                if w > widths.get(col, 0):
                    widths[col] = w
    for col, w in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = min(_MAX_COL_WIDTH, max(_MIN_COL_WIDTH, w))


def _finalize(ws: Worksheet, header_row: int, last_row: int, ncols: int) -> None:
    """フリーズペインとオートフィルタを設定する（プライマリ表に対して1回）。"""
    ncols = max(ncols, 1)
    ws.freeze_panes = f"A{header_row + 1}"
    last_row = max(last_row, header_row)
    ws.auto_filter.ref = f"A{header_row}:{get_column_letter(ncols)}{last_row}"
    ws.sheet_view.showGridLines = False


# ---------------------------------------------------------------------------
# inventory アクセスの共通ヘルパー（questions.py と同じ流儀だが独立実装）
# ---------------------------------------------------------------------------


def _sect(inv: dict, name: str) -> dict:
    val = inv.get(name)
    return val if isinstance(val, dict) else {}


def _lst(inv: dict, section: str, key: str) -> list:
    val = _sect(inv, section).get(key)
    return val if isinstance(val, list) else []


def _map(inv: dict, section: str, key: str) -> dict:
    val = _sect(inv, section).get(key)
    return val if isinstance(val, dict) else {}


def _dicts(inv: dict, section: str, key: str) -> list[dict]:
    return [x for x in _lst(inv, section, key) if isinstance(x, dict)]


def _tag(resource: dict, key: str = "Name", default: str = "") -> str:
    for t in (resource or {}).get("Tags") or []:
        if isinstance(t, dict) and t.get("Key") == key:
            return t.get("Value") or default
    return default


def _index_by(items: list[dict], key: str) -> dict[str, dict]:
    return {i.get(key): i for i in items if isinstance(i, dict) and i.get(key)}


_ENV_ORDER = ("stg3", "stg2", "check", "demo", "prod", "stg")


def _guess_env(name: Any) -> str:
    """リソース名から環境名を推定する（questions.py の `_guess_env` と同じ規則）。"""
    lowered = str(name or "").lower()
    for env in _ENV_ORDER:
        if env in lowered:
            return env
    if "production" in lowered:
        return "prod"
    if "staging" in lowered:
        return "stg"
    return "その他"


def _external_accounts(document: Any, account_id: str | None) -> set[str]:
    """AssumeRolePolicyDocument から自アカウント以外の12桁アカウントIDを拾う。"""
    if document is None:
        return set()
    text = document if isinstance(document, str) else str(document)
    found = set(re.findall(r"arn:aws[\w-]*:(?:iam|sts)::(\d{12}):", text))
    found |= set(re.findall(r"['\"]AWS['\"]\s*:\s*['\"](\d{12})['\"]", text))
    return {a for a in found if a and a != (account_id or "")}


def _parse_dt(value: Any) -> _dt.datetime | None:
    if not isinstance(value, str) or not value:
        return None
    text = value.strip().replace("Z", "+00:00")
    text = re.sub(r"([+-]\d{2})(\d{2})$", r"\1:\2", text)
    try:
        return _dt.datetime.fromisoformat(text)
    except ValueError:
        try:
            return _dt.datetime.fromisoformat(text[:19])
        except ValueError:
            return None


def _days_since(value: Any) -> int | None:
    parsed = _parse_dt(value)
    if parsed is None:
        return None
    today = _dt.datetime.now(parsed.tzinfo) if parsed.tzinfo else _dt.datetime.now()
    return (today.date() - parsed.date()).days


def _route_table_for_subnet(subnet_id: str, vpc_id: Any, route_tables: list[dict]) -> dict | None:
    for rt in route_tables:
        for assoc in rt.get("Associations") or []:
            if isinstance(assoc, dict) and assoc.get("SubnetId") == subnet_id:
                return rt
    for rt in route_tables:
        if rt.get("VpcId") != vpc_id:
            continue
        for assoc in rt.get("Associations") or []:
            if isinstance(assoc, dict) and assoc.get("Main"):
                return rt
    return None


def _is_public_route_table(rt: dict | None) -> bool:
    if not rt:
        return False
    for route in rt.get("Routes") or []:
        if not isinstance(route, dict):
            continue
        gw = str(route.get("GatewayId") or "")
        if gw.startswith("igw-"):
            return True
    return False


def _route_target(route: dict) -> str:
    for key in (
        "GatewayId", "NatGatewayId", "InstanceId", "NetworkInterfaceId",
        "VpcPeeringConnectionId", "TransitGatewayId", "LocalGatewayId",
        "CarrierGatewayId", "EgressOnlyInternetGatewayId",
    ):
        v = route.get(key)
        if v:
            return f"{v}"
    return "local" if route.get("Origin") == "CreateRouteTable" else "—"


def _route_destination(route: dict) -> str:
    for key in ("DestinationCidrBlock", "DestinationIpv6CidrBlock", "DestinationPrefixListId"):
        v = route.get(key)
        if v:
            return str(v)
    return "—"


# ---------------------------------------------------------------------------
# サマリ
# ---------------------------------------------------------------------------

#: (表示ラベル, セクション名, キー名) … サマリの「リソース種別ごとの件数」表
_COUNT_SPECS: tuple[tuple[str, str, str], ...] = (
    ("VPC", "network", "vpcs"),
    ("サブネット", "network", "subnets"),
    ("ルートテーブル", "network", "route_tables"),
    ("ネットワークACL", "network", "network_acls"),
    ("インターネットゲートウェイ", "network", "internet_gateways"),
    ("NATゲートウェイ", "network", "nat_gateways"),
    ("セキュリティグループ", "network", "security_groups"),
    ("ENI（ネットワークインターフェース）", "network", "network_interfaces"),
    ("VPCエンドポイント", "network", "vpc_endpoints"),
    ("Elastic IP", "network", "elastic_ips"),
    ("EC2インスタンス", "compute", "instances"),
    ("EBSボリューム", "compute", "volumes"),
    ("AMI", "compute", "images"),
    ("キーペア", "compute", "key_pairs"),
    ("SSM管理対象インスタンス", "compute", "ssm_managed_instances"),
    ("RDSインスタンス", "database", "db_instances"),
    ("RDSスナップショット", "database", "db_snapshots"),
    ("S3バケット", "storage", "buckets"),
    ("EFSファイルシステム", "storage", "efs_file_systems"),
    ("バックアッププラン", "storage", "backup_plans"),
    ("ロードバランサ（ALB/NLB）", "edge", "load_balancers"),
    ("ターゲットグループ", "edge", "target_groups"),
    ("ACM証明書", "edge", "acm_certificates"),
    ("CloudFrontディストリビューション", "edge", "cloudfront_distributions"),
    ("WAF Web ACL", "edge", "wafv2_web_acls"),
    ("Lambda関数", "serverless", "lambda_functions"),
    ("EventBridgeルール", "serverless", "eventbridge_rules"),
    ("DynamoDBテーブル", "serverless", "dynamodb_tables"),
    ("SQSキュー", "serverless", "sqs_queues"),
    ("SNSトピック", "serverless", "sns_topics"),
    ("CloudTrail証跡", "logging", "cloudtrail_trails"),
    ("Configレコーダ", "logging", "config_recorders"),
    ("VPCフローログ", "logging", "flow_logs"),
    ("CloudWatchロググループ", "logging", "cloudwatch_log_groups"),
    ("CloudWatchアラーム", "logging", "cloudwatch_alarms"),
    ("GuardDuty検出器", "security", "guardduty_detectors"),
    ("CloudFormationスタック", "security", "cloudformation_stacks"),
    ("CloudFormation StackSet", "security", "cloudformation_stack_sets"),
)


def _build_summary_sheet(
    ws: Worksheet,
    inv: dict,
    sheet_status: list[tuple[str, bool, int]],
) -> None:
    meta = _sect(inv, "meta")
    errors = inv.get("errors") if isinstance(inv.get("errors"), list) else []
    iam = _map(inv, "security", "iam")

    row = 1
    _write_title_row(ws, row, "収集メタ情報", 2)
    row += 1
    meta_headers = ["項目", "値"]
    meta_rows = [
        ["収集日時", meta.get("collected_at") or "—"],
        ["アカウント別名", meta.get("account_alias") or "—"],
        ["アカウントID", meta.get("account_id") or "—"],
        ["リージョン", meta.get("region") or "—"],
        ["awsprobe バージョン", meta.get("awsprobe_version") or "—"],
        ["マスキング", "有効" if meta.get("redacted") else "無効（生値を含む）"],
        ["実行したコレクタ", ", ".join(meta.get("collectors_run") or []) or "—"],
        ["収集エラー件数", len(errors)],
        ["IAMユーザー数", len(iam.get("users") or [])],
        ["IAMロール数", len(iam.get("roles") or [])],
    ]
    row = _write_table(ws, row, meta_headers, meta_rows)
    meta_header_row = row - len(meta_rows) - 1
    meta_last_row = row - 1
    row += 1

    _write_title_row(ws, row, "リソース種別ごとの件数", 2)
    row += 1
    count_headers = ["リソース種別", "件数"]
    count_rows: list[list[Any]] = []
    for label, section, key in _COUNT_SPECS:
        count_rows.append([label, len(_lst(inv, section, key))])
    row = _write_table(ws, row, count_headers, count_rows)
    row += 1

    _write_title_row(ws, row, "シート作成状況", 3)
    row += 1
    status_headers = ["シート名", "状態", "行数"]
    status_rows = [
        [label, "データあり" if created else "データなし", rows_n if created else 0]
        for label, created, rows_n in sheet_status
    ]
    row = _write_table(ws, row, status_headers, status_rows)
    row += 1

    _autosize(
        ws,
        [
            (meta_headers, meta_rows),
            (count_headers, count_rows),
            (status_headers, status_rows),
        ],
    )
    _finalize(ws, meta_header_row, meta_last_row, len(meta_headers))
    ws.freeze_panes = None  # サマリは複数表を積むためフリーズは付けない


# ---------------------------------------------------------------------------
# VPC・サブネット
# ---------------------------------------------------------------------------


def _build_vpc_subnet_sheet(wb: Workbook, inv: dict) -> int | None:
    vpcs = _dicts(inv, "network", "vpcs")
    subnets = _dicts(inv, "network", "subnets")
    if not vpcs and not subnets:
        return None

    route_tables = _dicts(inv, "network", "route_tables")
    instances = _dicts(inv, "compute", "instances")
    db_instances = _dicts(inv, "database", "db_instances")
    mount_targets = _dicts(inv, "storage", "efs_mount_targets")
    load_balancers = _dicts(inv, "edge", "load_balancers")
    enis = _dicts(inv, "network", "network_interfaces")

    ws = wb.create_sheet("VPC・サブネット")
    row = 1

    vpc_headers = ["VPCID", "名前", "CIDR", "セカンダリCIDR", "デフォルトVPCか", "状態", "所有者"]
    vpc_rows = []
    for v in vpcs:
        secondary = [
            a.get("CidrBlock")
            for a in v.get("CidrBlockAssociationSet") or []
            if isinstance(a, dict) and a.get("CidrBlock") and a.get("CidrBlock") != v.get("CidrBlock")
        ]
        vpc_rows.append(
            [v.get("VpcId"), _tag(v), v.get("CidrBlock"), secondary, v.get("IsDefault"),
             v.get("State"), v.get("OwnerId")]
        )
    _write_title_row(ws, row, "VPC", len(vpc_headers))
    row += 1
    row = _write_table(ws, row, vpc_headers, vpc_rows)
    vpc_header_row = row - len(vpc_rows) - 1
    vpc_last_row = row - 1
    row += 1

    def _resources_in_subnet(subnet_id: str) -> list[str]:
        found: list[str] = []
        for i in instances:
            if i.get("SubnetId") == subnet_id:
                found.append(f"EC2:{_tag(i) or i.get('InstanceId')}")
        for db in db_instances:
            group = db.get("DBSubnetGroup") or {}
            for s in group.get("Subnets") or []:
                if isinstance(s, dict) and s.get("SubnetIdentifier") == subnet_id:
                    found.append(f"RDS:{db.get('DBInstanceIdentifier')}")
        for mt in mount_targets:
            if mt.get("SubnetId") == subnet_id:
                found.append(f"EFS-MT:{mt.get('FileSystemId')}")
        for lb in load_balancers:
            for az in lb.get("AvailabilityZones") or []:
                if isinstance(az, dict) and az.get("SubnetId") == subnet_id:
                    found.append(f"LB:{lb.get('LoadBalancerName')}")
        return found

    subnet_headers = [
        "サブネットID", "Nameタグ", "CIDR", "実AZ", "AZ-ID", "種別(public/private)",
        "ルートテーブル", "DefaultForAz", "利用可能IP数", "関連リソース数", "関連リソース",
    ]
    subnet_rows = []
    for s in subnets:
        subnet_id = s.get("SubnetId")
        rt = _route_table_for_subnet(subnet_id, s.get("VpcId"), route_tables)
        rt_label = (_tag(rt) or rt.get("RouteTableId")) if rt else "—"
        kind = "public" if _is_public_route_table(rt) else "private"
        resources = _resources_in_subnet(subnet_id) if subnet_id else []
        subnet_rows.append(
            [
                subnet_id, _tag(s), s.get("CidrBlock"), s.get("AvailabilityZone"),
                s.get("AvailabilityZoneId"), kind, rt_label, s.get("DefaultForAz"),
                s.get("AvailableIpAddressCount"), len(resources), resources,
            ]
        )
    _write_title_row(ws, row, "サブネット", len(subnet_headers))
    row += 1
    row = _write_table(ws, row, subnet_headers, subnet_rows)
    subnet_header_row = row - len(subnet_rows) - 1
    subnet_last_row = row - 1

    total_rows = len(vpc_rows) + len(subnet_rows)
    _autosize(ws, [(vpc_headers, vpc_rows), (subnet_headers, subnet_rows)])
    # サブネット表を主表としてフリーズ／オートフィルタを設定する
    _finalize(ws, subnet_header_row, subnet_last_row, len(subnet_headers))
    return total_rows


# ---------------------------------------------------------------------------
# ルートテーブル
# ---------------------------------------------------------------------------


def _build_route_table_sheet(wb: Workbook, inv: dict) -> int | None:
    route_tables = _dicts(inv, "network", "route_tables")
    if not route_tables:
        return None
    subnets_idx = _index_by(_dicts(inv, "network", "subnets"), "SubnetId")

    headers = [
        "ルートテーブルID", "名前", "VPCID", "宛先", "ターゲット", "状態", "Origin",
        "Mainテーブルか", "関連サブネット",
    ]
    rows: list[list[Any]] = []
    for rt in route_tables:
        assoc_subnets = []
        is_main = False
        for a in rt.get("Associations") or []:
            if not isinstance(a, dict):
                continue
            if a.get("Main"):
                is_main = True
            sid = a.get("SubnetId")
            if sid:
                sub = subnets_idx.get(sid)
                assoc_subnets.append(f"{_tag(sub) or sid} ({sid})" if sub else sid)
        routes = rt.get("Routes") or []
        if not routes:
            rows.append(
                [rt.get("RouteTableId"), _tag(rt), rt.get("VpcId"), "—", "—", "—", "—",
                 is_main, assoc_subnets]
            )
            continue
        for route in routes:
            if not isinstance(route, dict):
                continue
            rows.append(
                [
                    rt.get("RouteTableId"), _tag(rt), rt.get("VpcId"),
                    _route_destination(route), _route_target(route),
                    route.get("State"), route.get("Origin"), is_main, assoc_subnets,
                ]
            )

    ws = wb.create_sheet("ルートテーブル")
    end_row = _write_table(ws, 1, headers, rows)
    _autosize(ws, [(headers, rows)])
    _finalize(ws, 1, end_row - 1, len(headers))
    return len(rows)


# ---------------------------------------------------------------------------
# セキュリティグループ
# ---------------------------------------------------------------------------


def _proto_label(proto: Any) -> str:
    p = str(proto)
    return {"-1": "全プロトコル", "6": "tcp", "17": "udp", "1": "icmp"}.get(p, p)


def _port_label(perm: dict) -> str:
    proto = perm.get("IpProtocol")
    if str(proto) == "-1":
        return "全ポート"
    fp, tp = perm.get("FromPort"), perm.get("ToPort")
    if fp is None and tp is None:
        return "全ポート"
    if fp == tp:
        return str(fp)
    return f"{fp}-{tp}"


def _build_security_group_sheet(wb: Workbook, inv: dict) -> int | None:
    sgs = _dicts(inv, "network", "security_groups")
    if not sgs:
        return None
    sg_idx = _index_by(sgs, "GroupId")

    headers = [
        "SG ID", "SG名", "VPCID", "方向", "プロトコル", "ポート", "送信元種別",
        "送信元", "説明", "全世界公開(0.0.0.0/0)",
    ]
    rows: list[list[Any]] = []
    for sg in sgs:
        for direction, key in (("インバウンド", "IpPermissions"), ("アウトバウンド", "IpPermissionsEgress")):
            for perm in sg.get(key) or []:
                if not isinstance(perm, dict):
                    continue
                proto_label = _proto_label(perm.get("IpProtocol"))
                port_label = _port_label(perm)
                entries: list[tuple[str, str, str]] = []
                for r in perm.get("IpRanges") or []:
                    if isinstance(r, dict) and r.get("CidrIp"):
                        entries.append(("CIDR", r["CidrIp"], r.get("Description") or ""))
                for r in perm.get("Ipv6Ranges") or []:
                    if isinstance(r, dict) and r.get("CidrIpv6"):
                        entries.append(("CIDR(v6)", r["CidrIpv6"], r.get("Description") or ""))
                for p in perm.get("UserIdGroupPairs") or []:
                    if not isinstance(p, dict):
                        continue
                    gid = p.get("GroupId")
                    ref_sg = sg_idx.get(gid)
                    label = f"{ref_sg.get('GroupName')} ({gid})" if ref_sg else gid
                    entries.append(("SG参照", label, p.get("Description") or ""))
                for p in perm.get("PrefixListIds") or []:
                    if isinstance(p, dict) and p.get("PrefixListId"):
                        entries.append(("プレフィックスリスト", p["PrefixListId"], p.get("Description") or ""))
                if not entries:
                    entries = [("—", "（送信元の記載なし）", "")]
                for source_type, source, desc in entries:
                    is_world = source.startswith("0.0.0.0/0") or source.startswith("::/0")
                    rows.append(
                        [
                            sg.get("GroupId"), sg.get("GroupName"), sg.get("VpcId"),
                            direction, proto_label, port_label, source_type, source,
                            desc or sg.get("Description"), is_world,
                        ]
                    )

    def _row_fill(values: Sequence[Any]) -> tuple[str, str] | None:
        # 「全世界公開」列（最終列）が真なら赤系で強調する
        return (_RED_FILL_RGB, _RED_FONT_RGB) if values[-1] else None

    ws = wb.create_sheet("セキュリティグループ")
    end_row = _write_table(ws, 1, headers, rows, row_fill=_row_fill)
    _autosize(ws, [(headers, rows)])
    _finalize(ws, 1, end_row - 1, len(headers))
    return len(rows)


# ---------------------------------------------------------------------------
# ネットワークACL
# ---------------------------------------------------------------------------


def _build_network_acl_sheet(wb: Workbook, inv: dict) -> int | None:
    nacls = _dicts(inv, "network", "network_acls")
    if not nacls:
        return None
    subnets_idx = _index_by(_dicts(inv, "network", "subnets"), "SubnetId")

    headers = [
        "NACL ID", "VPCID", "デフォルトか", "ルール番号", "方向", "プロトコル",
        "アクション", "送信元/宛先CIDR", "ポート範囲", "関連サブネット",
    ]
    rows: list[list[Any]] = []
    for nacl in nacls:
        assoc_subnets = []
        for a in nacl.get("Associations") or []:
            if isinstance(a, dict) and a.get("SubnetId"):
                sid = a["SubnetId"]
                sub = subnets_idx.get(sid)
                assoc_subnets.append(f"{_tag(sub) or sid} ({sid})" if sub else sid)
        entries = nacl.get("Entries") or []
        if not entries:
            rows.append(
                [nacl.get("NetworkAclId"), nacl.get("VpcId"), nacl.get("IsDefault"),
                 "—", "—", "—", "—", "—", "—", assoc_subnets]
            )
            continue
        for e in entries:
            if not isinstance(e, dict):
                continue
            port_range = e.get("PortRange") or {}
            port_text = (
                f"{port_range.get('From')}-{port_range.get('To')}" if port_range else "全ポート"
            )
            rows.append(
                [
                    nacl.get("NetworkAclId"), nacl.get("VpcId"), nacl.get("IsDefault"),
                    e.get("RuleNumber"),
                    "アウトバウンド" if e.get("Egress") else "インバウンド",
                    _proto_label(e.get("Protocol")),
                    "許可" if str(e.get("RuleAction")).lower() == "allow" else "拒否",
                    e.get("CidrBlock") or e.get("Ipv6CidrBlock") or "—",
                    port_text, assoc_subnets,
                ]
            )

    ws = wb.create_sheet("ネットワークACL")
    end_row = _write_table(ws, 1, headers, rows)
    _autosize(ws, [(headers, rows)])
    _finalize(ws, 1, end_row - 1, len(headers))
    return len(rows)


# ---------------------------------------------------------------------------
# EC2
# ---------------------------------------------------------------------------


def _build_ec2_sheet(wb: Workbook, inv: dict) -> int | None:
    instances = _dicts(inv, "compute", "instances")
    if not instances:
        return None
    managed_ids = {
        m.get("InstanceId") for m in _dicts(inv, "compute", "ssm_managed_instances") if m.get("InstanceId")
    }
    subnets_idx = _index_by(_dicts(inv, "network", "subnets"), "SubnetId")

    headers = [
        "インスタンス名", "インスタンスID", "タイプ", "AZ", "サブネット", "プライベートIP",
        "パブリックIPの有無", "AMI", "キーペア名", "IAMロール", "IMDSv2必須か",
        "SSM管理下か", "状態",
    ]
    rows = []
    for i in instances:
        placement = i.get("Placement") or {}
        state = i.get("State") or {}
        subnet_id = i.get("SubnetId")
        sub = subnets_idx.get(subnet_id)
        subnet_label = f"{_tag(sub) or subnet_id} ({subnet_id})" if sub else (subnet_id or "—")
        profile = i.get("IamInstanceProfile") or {}
        role_arn = profile.get("Arn") or ""
        role_name = role_arn.rsplit("/", 1)[-1] if role_arn else "—"
        metadata = i.get("MetadataOptions") or {}
        http_tokens = metadata.get("HttpTokens")
        imds_required = "必須(required)" if http_tokens == "required" else (
            "任意(optional)" if http_tokens else "不明"
        )
        rows.append(
            [
                _tag(i), i.get("InstanceId"), i.get("InstanceType"),
                placement.get("AvailabilityZone"), subnet_label, i.get("PrivateIpAddress"),
                "あり" if i.get("PublicIpAddress") else "なし", i.get("ImageId"),
                i.get("KeyName") or "—", role_name, imds_required,
                (i.get("InstanceId") in managed_ids), state.get("Name"),
            ]
        )

    ws = wb.create_sheet("EC2")
    end_row = _write_table(ws, 1, headers, rows)
    _autosize(ws, [(headers, rows)])
    _finalize(ws, 1, end_row - 1, len(headers))
    return len(rows)


# ---------------------------------------------------------------------------
# RDS
# ---------------------------------------------------------------------------


def _build_rds_sheet(wb: Workbook, inv: dict) -> int | None:
    db_instances = _dicts(inv, "database", "db_instances")
    if not db_instances:
        return None
    sg_idx = _index_by(_dicts(inv, "network", "security_groups"), "GroupId")

    headers = [
        "識別子", "エンジン", "バージョン", "クラス", "MultiAZ", "バックアップ保持日数",
        "暗号化", "セキュリティグループ", "サブネットグループ", "パラメータグループ",
        "公開設定(PubliclyAccessible)", "削除保護",
    ]
    rows = []
    for db in db_instances:
        sg_labels = []
        for g in db.get("VpcSecurityGroups") or []:
            if not isinstance(g, dict):
                continue
            gid = g.get("VpcSecurityGroupId")
            ref = sg_idx.get(gid)
            sg_labels.append(f"{ref.get('GroupName')} ({gid})" if ref else gid)
        param_groups = [
            p.get("DBParameterGroupName")
            for p in db.get("DBParameterGroups") or []
            if isinstance(p, dict) and p.get("DBParameterGroupName")
        ]
        subnet_group = (db.get("DBSubnetGroup") or {}).get("DBSubnetGroupName")
        rows.append(
            [
                db.get("DBInstanceIdentifier"), db.get("Engine"), db.get("EngineVersion"),
                db.get("DBInstanceClass"), db.get("MultiAZ"), db.get("BackupRetentionPeriod"),
                db.get("StorageEncrypted"), sg_labels, subnet_group, param_groups,
                db.get("PubliclyAccessible"), db.get("DeletionProtection"),
            ]
        )

    ws = wb.create_sheet("RDS")
    end_row = _write_table(ws, 1, headers, rows)
    _autosize(ws, [(headers, rows)])
    _finalize(ws, 1, end_row - 1, len(headers))
    return len(rows)


# ---------------------------------------------------------------------------
# ロードバランサ
# ---------------------------------------------------------------------------


def _lb_attr(lb: dict, key: str) -> str | None:
    for a in lb.get("Attributes") or []:
        if isinstance(a, dict) and a.get("Key") == key:
            return a.get("Value")
    return None


def _build_load_balancer_sheet(wb: Workbook, inv: dict) -> int | None:
    lbs = _dicts(inv, "edge", "load_balancers")
    if not lbs:
        return None
    listeners = _dicts(inv, "edge", "listeners")
    target_groups = _dicts(inv, "edge", "target_groups")
    web_acls = _dicts(inv, "edge", "wafv2_web_acls")

    listeners_by_lb: dict[str, list[dict]] = {}
    for l in listeners:
        listeners_by_lb.setdefault(l.get("LoadBalancerArn"), []).append(l)

    tg_by_lb: dict[str, list[dict]] = {}
    for tg in target_groups:
        for arn in tg.get("LoadBalancerArns") or []:
            tg_by_lb.setdefault(arn, []).append(tg)

    waf_by_resource: dict[str, str] = {}
    for acl in web_acls:
        for arn in acl.get("AssociatedResourceArns") or []:
            waf_by_resource[arn] = acl.get("Name")

    headers = [
        "名前", "種別", "スキーム", "VPCID", "AZ", "アクセスログ有効/無効",
        "WAF関連付け", "リスナーポート/プロトコル", "TLSポリシー",
        "ターゲットグループ", "対象ターゲット数",
    ]
    rows = []
    for lb in lbs:
        arn = lb.get("LoadBalancerArn")
        access_log_enabled = _lb_attr(lb, "access_logs.s3.enabled")
        azs = [az.get("ZoneName") for az in lb.get("AvailabilityZones") or [] if isinstance(az, dict)]
        tgs = tg_by_lb.get(arn, [])
        tg_labels = [tg.get("TargetGroupName") for tg in tgs]
        target_count = sum(len(tg.get("Targets") or []) for tg in tgs)
        waf = waf_by_resource.get(arn) or "なし"
        lb_listeners = listeners_by_lb.get(arn, [])
        if not lb_listeners:
            rows.append(
                [
                    lb.get("LoadBalancerName"), lb.get("Type"), lb.get("Scheme"), lb.get("VpcId"),
                    azs, access_log_enabled, waf, "—", "—", tg_labels, target_count,
                ]
            )
            continue
        for l in lb_listeners:
            rows.append(
                [
                    lb.get("LoadBalancerName"), lb.get("Type"), lb.get("Scheme"), lb.get("VpcId"),
                    azs, access_log_enabled, waf,
                    f"{l.get('Port')}/{l.get('Protocol')}", l.get("SslPolicy") or "—",
                    tg_labels, target_count,
                ]
            )

    ws = wb.create_sheet("ロードバランサ")
    end_row = _write_table(ws, 1, headers, rows)
    _autosize(ws, [(headers, rows)])
    _finalize(ws, 1, end_row - 1, len(headers))
    return len(rows)


# ---------------------------------------------------------------------------
# S3
# ---------------------------------------------------------------------------


def _s3_encryption_label(enc: Any) -> str:
    if not isinstance(enc, dict):
        return "無効"
    rules = enc.get("Rules") or []
    for r in rules:
        if not isinstance(r, dict):
            continue
        algo = ((r.get("ApplyServerSideEncryptionByDefault") or {}).get("SSEAlgorithm"))
        if algo:
            return algo
    return "有効(詳細不明)" if rules else "無効"


def _build_s3_sheet(wb: Workbook, inv: dict) -> int | None:
    buckets = _dicts(inv, "storage", "buckets")
    if not buckets:
        return None

    headers = [
        "バケット名", "リージョン", "暗号化", "バージョニング",
        "PAB:BlockPublicAcls", "PAB:IgnorePublicAcls", "PAB:BlockPublicPolicy",
        "PAB:RestrictPublicBuckets", "ポリシーによる公開", "静的ホスティング", "ライフサイクル有無",
    ]
    rows = []
    for b in buckets:
        pab = b.get("PublicAccessBlock") if isinstance(b.get("PublicAccessBlock"), dict) else {}
        versioning = b.get("Versioning") if isinstance(b.get("Versioning"), dict) else {}
        policy_status = b.get("PolicyStatus") if isinstance(b.get("PolicyStatus"), dict) else {}
        is_public = policy_status.get("IsPublic")
        lifecycle = b.get("Lifecycle")
        rows.append(
            [
                b.get("Name"), b.get("Region"), _s3_encryption_label(b.get("Encryption")),
                versioning.get("Status") or "未設定",
                pab.get("BlockPublicAcls"), pab.get("IgnorePublicAcls"),
                pab.get("BlockPublicPolicy"), pab.get("RestrictPublicBuckets"),
                is_public if is_public is not None else "不明",
                "有効" if b.get("Website") else "無効",
                "あり" if lifecycle else "なし",
            ]
        )

    ws = wb.create_sheet("S3")
    end_row = _write_table(ws, 1, headers, rows)
    _autosize(ws, [(headers, rows)])
    _finalize(ws, 1, end_row - 1, len(headers))
    return len(rows)


# ---------------------------------------------------------------------------
# EFS
# ---------------------------------------------------------------------------


def _fmt_bytes(value: Any) -> str:
    try:
        n = float(value)
    except (TypeError, ValueError):
        return "—"
    for unit in ("B", "KB", "MB", "GB", "TB"):
        if n < 1024 or unit == "TB":
            return f"{n:.1f}{unit}"
        n /= 1024
    return f"{n:.1f}TB"


def _build_efs_sheet(wb: Workbook, inv: dict) -> int | None:
    file_systems = _dicts(inv, "storage", "efs_file_systems")
    if not file_systems:
        return None
    mount_targets = _dicts(inv, "storage", "efs_mount_targets")
    access_points = _dicts(inv, "storage", "efs_access_points")
    backup_policies = _map(inv, "storage", "efs_backup_policies")

    mt_by_fs: dict[str, list[dict]] = {}
    for mt in mount_targets:
        mt_by_fs.setdefault(mt.get("FileSystemId"), []).append(mt)
    ap_by_fs: dict[str, list[dict]] = {}
    for ap in access_points:
        ap_by_fs.setdefault(ap.get("FileSystemId"), []).append(ap)

    headers = [
        "ファイルシステムID", "名前", "スループットモード", "暗号化", "サイズ",
        "マウントターゲット(AZ・サブネット)", "アクセスポイントのパス", "バックアップポリシー",
    ]
    rows = []
    for fs in file_systems:
        fs_id = fs.get("FileSystemId")
        mts = mt_by_fs.get(fs_id, [])
        mt_labels = [f"{m.get('AvailabilityZoneName')}:{m.get('SubnetId')}" for m in mts]
        aps = ap_by_fs.get(fs_id, [])
        ap_paths = [
            (a.get("RootDirectory") or {}).get("Path") for a in aps
            if isinstance(a.get("RootDirectory"), dict) and (a.get("RootDirectory") or {}).get("Path")
        ]
        backup = backup_policies.get(fs_id) if isinstance(backup_policies.get(fs_id), dict) else {}
        size = (fs.get("SizeInBytes") or {}).get("Value") if isinstance(fs.get("SizeInBytes"), dict) else None
        rows.append(
            [
                fs_id, fs.get("Name") or _tag(fs), fs.get("ThroughputMode"), fs.get("Encrypted"),
                _fmt_bytes(size), mt_labels, ap_paths, backup.get("Status") or "不明",
            ]
        )

    ws = wb.create_sheet("EFS")
    end_row = _write_table(ws, 1, headers, rows)
    _autosize(ws, [(headers, rows)])
    _finalize(ws, 1, end_row - 1, len(headers))
    return len(rows)


# ---------------------------------------------------------------------------
# Lambda・EventBridge
# ---------------------------------------------------------------------------


def _build_lambda_eventbridge_sheet(wb: Workbook, inv: dict) -> int | None:
    functions = _dicts(inv, "serverless", "lambda_functions")
    rules = _dicts(inv, "serverless", "eventbridge_rules")
    if not functions and not rules:
        return None

    ws = wb.create_sheet("Lambda・EventBridge")
    row = 1
    total = 0

    lambda_headers = ["関数名", "ランタイム", "最終更新", "トリガ", "関数URL", "実行ロール"]
    lambda_rows = []
    for fn in functions:
        triggers = [
            m.get("EventSourceArn") for m in fn.get("EventSourceMappings") or []
            if isinstance(m, dict) and m.get("EventSourceArn")
        ]
        for r in rules:
            for t in r.get("Targets") or []:
                if isinstance(t, dict) and t.get("Arn") == fn.get("FunctionArn"):
                    triggers.append(f"EventBridge:{r.get('Name')}")
        url_config = fn.get("UrlConfig")
        url_label = "あり(" + str((url_config or {}).get("AuthType")) + ")" if url_config else "なし"
        role_arn = fn.get("Role") or ""
        lambda_rows.append(
            [
                fn.get("FunctionName"), fn.get("Runtime"), fn.get("LastModified"),
                triggers or "（トリガ不明・要調査）", url_label, role_arn.rsplit("/", 1)[-1] if role_arn else "—",
            ]
        )
    if lambda_rows:
        _write_title_row(ws, row, "Lambda 関数", len(lambda_headers))
        row += 1
        row = _write_table(ws, row, lambda_headers, lambda_rows)
        lambda_header_row = row - len(lambda_rows) - 1
        lambda_last_row = row - 1
        row += 1
        total += len(lambda_rows)
    else:
        lambda_header_row = lambda_last_row = None

    eb_headers = ["ルール名", "イベントバス", "スケジュール/パターン", "状態", "ターゲット"]
    eb_rows = []
    for r in rules:
        targets = [t.get("Arn") or t.get("Id") for t in r.get("Targets") or [] if isinstance(t, dict)]
        schedule = r.get("ScheduleExpression") or (r.get("EventPattern") and "イベントパターン指定") or "—"
        eb_rows.append([r.get("Name"), r.get("EventBusName"), schedule, r.get("State"), targets])
    if eb_rows:
        _write_title_row(ws, row, "EventBridge ルール", len(eb_headers))
        row += 1
        row = _write_table(ws, row, eb_headers, eb_rows)
        eb_header_row = row - len(eb_rows) - 1
        eb_last_row = row - 1
        total += len(eb_rows)
    else:
        eb_header_row = eb_last_row = None

    _autosize(ws, [(lambda_headers, lambda_rows), (eb_headers, eb_rows)])
    # 行数の多い方を主表としてフリーズ／オートフィルタを設定する
    if len(eb_rows) >= len(lambda_rows) and eb_header_row is not None:
        _finalize(ws, eb_header_row, eb_last_row, len(eb_headers))
    elif lambda_header_row is not None:
        _finalize(ws, lambda_header_row, lambda_last_row, len(lambda_headers))
    return total


# ---------------------------------------------------------------------------
# ログ・証跡
# ---------------------------------------------------------------------------


def _build_logging_sheet(wb: Workbook, inv: dict) -> int | None:
    trails = _dicts(inv, "logging", "cloudtrail_trails")
    recorders = _dicts(inv, "logging", "config_recorders")
    flow_logs = _dicts(inv, "logging", "flow_logs")
    log_groups = _dicts(inv, "logging", "cloudwatch_log_groups")
    if not (trails or recorders or flow_logs or log_groups):
        return None

    log_group_kms = _map(inv, "logging", "log_group_kms")
    log_group_retention = {
        g.get("logGroupName"): g.get("retentionInDays")
        for g in log_groups
        if g.get("logGroupName")
    }

    ws = wb.create_sheet("ログ・証跡")
    row = 1
    total = 0
    header_rows: list[tuple[int, int, int]] = []  # (header_row, last_row, ncols)

    if trails:
        headers = [
            "証跡名", "S3バケット", "全リージョン", "記録中か(IsLogging)", "改ざん検知",
            "KMS暗号化", "保持期間",
        ]
        rows = [
            [
                t.get("Name"), t.get("S3BucketName"), t.get("IsMultiRegionTrail"),
                t.get("IsLogging"), t.get("LogFileValidationEnabled"),
                "あり" if t.get("KmsKeyId") else "なし（SSE-S3）",
                "—（配信先S3のライフサイクルに依存）",
            ]
            for t in trails
        ]
        _write_title_row(ws, row, "CloudTrail 証跡", len(headers))
        row += 1
        row = _write_table(ws, row, headers, rows)
        header_rows.append((row - len(rows) - 1, row - 1, len(headers)))
        total += len(rows)
        row += 1
        _autosize(ws, [(headers, rows)])

    if recorders:
        headers = ["レコーダ名", "ロールARN", "記録中か", "最終ステータス", "保持期間"]
        rows = []
        for r in recorders:
            status = r.get("Status") or {}
            rows.append(
                [
                    r.get("name"), r.get("roleARN"), status.get("recording"),
                    status.get("lastStatus"), "—（継続記録・保持期間の概念なし）",
                ]
            )
        _write_title_row(ws, row, "AWS Config レコーダ", len(headers))
        row += 1
        row = _write_table(ws, row, headers, rows)
        header_rows.append((row - len(rows) - 1, row - 1, len(headers)))
        total += len(rows)
        row += 1
        _autosize(ws, [(headers, rows)])

    if flow_logs:
        headers = ["FlowLogID", "対象リソース", "配信先種別", "配信先", "状態", "保持期間"]
        rows = []
        for f in flow_logs:
            dest = f.get("LogDestination") or ""
            dest_type = f.get("LogDestinationType")
            if dest_type == "cloud-watch-logs":
                group_name = dest.rsplit(":", 1)[-1] if dest else None
                retention = log_group_retention.get(group_name)
                retention_label = f"{retention}日" if retention else "無期限保持"
            else:
                retention_label = "—（配信先S3のライフサイクルに依存）"
            rows.append(
                [f.get("FlowLogId"), f.get("ResourceId"), dest_type, dest,
                 f.get("FlowLogStatus"), retention_label]
            )
        _write_title_row(ws, row, "VPC フローログ", len(headers))
        row += 1
        row = _write_table(ws, row, headers, rows)
        header_rows.append((row - len(rows) - 1, row - 1, len(headers)))
        total += len(rows)
        row += 1
        _autosize(ws, [(headers, rows)])

    if log_groups:
        headers = ["ロググループ名", "保持期間", "KMS暗号化"]
        rows = []
        for g in log_groups:
            name = g.get("logGroupName")
            retention = g.get("retentionInDays")
            kms = log_group_kms.get(name)
            rows.append(
                [name, f"{retention}日" if retention else "無期限保持", "あり" if kms else "なし"]
            )
        _write_title_row(ws, row, "CloudWatch ロググループ", len(headers))
        row += 1
        row = _write_table(ws, row, headers, rows)
        header_rows.append((row - len(rows) - 1, row - 1, len(headers)))
        total += len(rows)
        row += 1
        _autosize(ws, [(headers, rows)])

    if header_rows:
        # 最も行数の多いブロックを主表としてフリーズ／オートフィルタを設定する
        primary = max(header_rows, key=lambda t: t[1] - t[0])
        _finalize(ws, primary[0], primary[1], primary[2])
    return total


# ---------------------------------------------------------------------------
# IAM
# ---------------------------------------------------------------------------


def _build_iam_sheet(wb: Workbook, inv: dict) -> int | None:
    iam = _map(inv, "security", "iam")
    roles = [r for r in (iam.get("roles") or []) if isinstance(r, dict)]
    users = [u for u in (iam.get("users") or []) if isinstance(u, dict)]
    access_keys = [k for k in (iam.get("access_keys") or []) if isinstance(k, dict)]
    if not roles and not users and not access_keys:
        return None

    account_id = _sect(inv, "meta").get("account_id")
    attached = [a for a in (iam.get("policies_attached_summary") or []) if isinstance(a, dict)]
    attached_count: dict[tuple[str, str], int] = {}
    for a in attached:
        key = (a.get("PrincipalType"), a.get("Name"))
        attached_count[key] = len(a.get("AttachedPolicies") or [])
    mfa = iam.get("mfa_devices") if isinstance(iam.get("mfa_devices"), dict) else {}
    mfa_users = mfa.get("UserDevices") if isinstance(mfa.get("UserDevices"), dict) else {}
    keys_by_user: dict[str, int] = {}
    for k in access_keys:
        keys_by_user[k.get("UserName")] = keys_by_user.get(k.get("UserName"), 0) + 1

    ws = wb.create_sheet("IAM")
    row = 1
    total = 0
    header_rows: list[tuple[int, int, int]] = []

    if roles:
        headers = ["ロール名", "Path", "作成日", "外部アカウント信頼", "信頼している外部アカウントID", "付与ポリシー数"]
        rows = []

        def _role_fill(values: Sequence[Any]) -> tuple[str, str] | None:
            return (_RED_FILL_RGB, _RED_FONT_RGB) if values[3] else None

        for r in roles:
            ext = _external_accounts(r.get("AssumeRolePolicyDocument"), account_id)
            name = r.get("RoleName")
            rows.append(
                [
                    name, r.get("Path"), r.get("CreateDate"), bool(ext), sorted(ext),
                    attached_count.get(("role", name), 0),
                ]
            )
        _write_title_row(ws, row, "IAMロール（外部アカウントを信頼するものは赤字で強調）", len(headers))
        row += 1
        row = _write_table(ws, row, headers, rows, row_fill=_role_fill)
        header_rows.append((row - len(rows) - 1, row - 1, len(headers)))
        total += len(rows)
        row += 1
        _autosize(ws, [(headers, rows)])

    if users:
        headers = ["ユーザー名", "Path", "作成日", "最終パスワード利用", "MFA登録", "アクセスキー数", "付与ポリシー数"]
        rows = []
        for u in users:
            name = u.get("UserName")
            rows.append(
                [
                    name, u.get("Path"), u.get("CreateDate"), u.get("PasswordLastUsed") or "—",
                    bool(mfa_users.get(name)), keys_by_user.get(name, 0),
                    attached_count.get(("user", name), 0),
                ]
            )
        _write_title_row(ws, row, "IAMユーザー", len(headers))
        row += 1
        row = _write_table(ws, row, headers, rows)
        header_rows.append((row - len(rows) - 1, row - 1, len(headers)))
        total += len(rows)
        row += 1
        _autosize(ws, [(headers, rows)])

    if access_keys:
        headers = ["ユーザー名", "アクセスキーID(末尾4桁)", "状態", "作成日", "経過日数", "最終使用日", "最終使用サービス", "最終使用リージョン"]
        rows = []
        for k in access_keys:
            rows.append(
                [
                    k.get("UserName"), k.get("AccessKeyId"), k.get("Status"), k.get("CreateDate"),
                    _days_since(k.get("CreateDate")), k.get("LastUsedDate") or "未使用",
                    k.get("ServiceName") or "—", k.get("Region") or "—",
                ]
            )
        _write_title_row(ws, row, "アクセスキーの経過日数", len(headers))
        row += 1
        row = _write_table(ws, row, headers, rows)
        header_rows.append((row - len(rows) - 1, row - 1, len(headers)))
        total += len(rows)
        row += 1
        _autosize(ws, [(headers, rows)])

    if header_rows:
        primary = max(header_rows, key=lambda t: t[1] - t[0])
        _finalize(ws, primary[0], primary[1], primary[2])
    return total


# ---------------------------------------------------------------------------
# セキュリティ設定評価（posture_results）
# ---------------------------------------------------------------------------


def _attr(obj: Any, name: str, default: Any = None) -> Any:
    if obj is None:
        return default
    if isinstance(obj, dict):
        return obj.get(name, default)
    return getattr(obj, name, default)


def _check_field(result: Any, name: str, default: Any = None) -> Any:
    """CheckResult（または to_dict() の結果）から `check.<name>` を取り出す。"""
    check = _attr(result, "check")
    if check is not None:
        return _attr(check, name, default)
    return _attr(result, name, default)


_POSTURE_ROW_STYLE = {
    "未実施": (_RED_FILL_RGB, _RED_FONT_RGB),
    "一部実施": (_YELLOW_FILL_RGB, _YELLOW_FONT_RGB),
}


def _build_posture_sheet(wb: Workbook, posture_results: Any) -> int | None:
    results = list(posture_results or [])
    if not results:
        return None

    headers = [
        "チェックID", "ドメイン", "チェック名", "深刻度", "実施状況", "結論",
        "是正方針", "満たしたリソース数", "満たさなかったリソース数", "根拠",
    ]
    rows = []
    for r in results:
        status = _attr(r, "status")
        rows.append(
            [
                _check_field(r, "cid"), _check_field(r, "domain"), _check_field(r, "title"),
                _check_field(r, "severity"), status, _attr(r, "summary"),
                _attr(r, "remediation") or "—",
                len(_attr(r, "passed") or []), len(_attr(r, "failed") or []),
                _attr(r, "evidence") or [],
            ]
        )

    def _row_fill(values: Sequence[Any]) -> tuple[str, str] | None:
        return _POSTURE_ROW_STYLE.get(str(values[4]))

    ws = wb.create_sheet("セキュリティ設定評価")
    end_row = _write_table(ws, 1, headers, rows, row_fill=_row_fill)
    _autosize(ws, [(headers, rows)])
    _finalize(ws, 1, end_row - 1, len(headers))
    return len(rows)


# ---------------------------------------------------------------------------
# 未確認事項の判定（answers）
# ---------------------------------------------------------------------------


def _build_answers_sheet(wb: Workbook, answers: Any) -> int | None:
    items = list(answers or [])
    if not items:
        return None

    headers = ["Q番号", "カテゴリ", "設問", "判定", "結論", "根拠", "確認手順"]
    rows = []
    for a in items:
        rows.append(
            [
                _attr(a, "qid"), _attr(a, "category"), _attr(a, "title"), _attr(a, "status"),
                _attr(a, "summary"), _attr(a, "evidence") or [], _attr(a, "manual_steps") or "—",
            ]
        )

    ws = wb.create_sheet("未確認事項の判定")
    end_row = _write_table(ws, 1, headers, rows)
    _autosize(ws, [(headers, rows)])
    _finalize(ws, 1, end_row - 1, len(headers))
    return len(rows)


# ---------------------------------------------------------------------------
# 収集エラー
# ---------------------------------------------------------------------------


def _build_errors_sheet(wb: Workbook, inv: dict) -> int | None:
    errors = inv.get("errors") if isinstance(inv.get("errors"), list) else []
    errors = [e for e in errors if isinstance(e, dict)]
    if not errors:
        return None

    headers = ["サービス", "オペレーション", "コード", "コンテキスト", "メッセージ"]
    rows = [
        [e.get("service"), e.get("operation"), e.get("code"), e.get("context") or "—", e.get("message")]
        for e in errors
    ]

    ws = wb.create_sheet("収集エラー")
    end_row = _write_table(ws, 1, headers, rows)
    _autosize(ws, [(headers, rows)])
    _finalize(ws, 1, end_row - 1, len(headers))
    return len(rows)


# ---------------------------------------------------------------------------
# エントリポイント
# ---------------------------------------------------------------------------


def build_workbook(
    inventory: dict,
    out_path: str,
    *,
    posture_results: Any = None,
    answers: Any = None,
) -> str:
    """inventory dict から Excel 棚卸し表を作り、`out_path` に保存してそのパスを返す。

    Args:
        inventory: `awsprobe collect` が出力した inventory dict。
        out_path: 保存先パス（.xlsx）。
        posture_results: `posture.py` の判定結果（`CheckResult` の list、または
            それぞれの `to_dict()` の list）。無ければ「セキュリティ設定評価」
            シートは作らない。
        answers: `questions.py` の判定結果（`Answer` の list、または
            それぞれの `to_dict()` の list）。無ければ「未確認事項の判定」
            シートは作らない。
    """
    inv = inventory if isinstance(inventory, dict) else {}

    wb = Workbook()
    wb.remove(wb.active)
    ws_summary = wb.create_sheet("サマリ")

    sheet_status: list[tuple[str, bool, int]] = []

    def _add(label: str, rows: int | None) -> None:
        sheet_status.append((label, rows is not None, rows or 0))

    _add("VPC・サブネット", _build_vpc_subnet_sheet(wb, inv))
    _add("ルートテーブル", _build_route_table_sheet(wb, inv))
    _add("セキュリティグループ", _build_security_group_sheet(wb, inv))
    _add("ネットワークACL", _build_network_acl_sheet(wb, inv))
    _add("EC2", _build_ec2_sheet(wb, inv))
    _add("RDS", _build_rds_sheet(wb, inv))
    _add("ロードバランサ", _build_load_balancer_sheet(wb, inv))
    _add("S3", _build_s3_sheet(wb, inv))
    _add("EFS", _build_efs_sheet(wb, inv))
    _add("Lambda・EventBridge", _build_lambda_eventbridge_sheet(wb, inv))
    _add("ログ・証跡", _build_logging_sheet(wb, inv))
    _add("IAM", _build_iam_sheet(wb, inv))
    _add("セキュリティ設定評価", _build_posture_sheet(wb, posture_results))
    _add("未確認事項の判定", _build_answers_sheet(wb, answers))
    _add("収集エラー", _build_errors_sheet(wb, inv))

    _build_summary_sheet(ws_summary, inv, sheet_status)

    wb.save(out_path)
    return out_path
